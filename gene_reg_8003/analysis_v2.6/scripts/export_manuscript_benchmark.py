"""Export the published 415-pair workbook to a portable, auditable CSV.

Run from any directory: python gene_reg_8003/analysis_v2.6/scripts/export_manuscript_benchmark.py
Requires openpyxl. The source workbook and its partition labels are preserved.
"""
import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "analysis_v2.6/results/Supplementary_Tables_S2_S5.xlsx"
OUTPUT = ROOT / "analysis_v2.6/manuscript_benchmark_415.csv"
FIELDS = ["record_id", "source", "subset", "partition", "label", "tf", "gene", "ensembl_id",
          "chip_score", "motif_score", "literature_score", "weighted_score",
          "weighted_score_normalized", "prediction_final"]


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = next(sheet for sheet in workbook if "415" in sheet.title)
    rows = []
    for values in sheet.iter_rows(min_row=4, values_only=True):
        if not values[1]:
            continue
        row = dict(zip(FIELDS, values[1:15]))
        row["label"] = int(row["label"] == "Positive")
        rows.append(row)
    assert len(rows) == len({(row["tf"], row["gene"]) for row in rows}) == 415
    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"Exported {len(rows)} published rows to {OUTPUT}")


if __name__ == "__main__":
    main()
